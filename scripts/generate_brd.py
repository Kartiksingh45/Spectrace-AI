"""Generates Spectrace_AI_BRD.docx — Business Requirements Document for Spectrace AI."""
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import datetime

NAVY = RGBColor(0x1B, 0x2A, 0x4A)
GREY = RGBColor(0x44, 0x44, 0x44)

doc = Document()

# ---------- base style ----------
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(10.5)

for i in range(1, 4):
    hstyle = doc.styles[f"Heading {i}"]
    hstyle.font.name = "Calibri"
    hstyle.font.color.rgb = NAVY
    hstyle.font.bold = True

doc.styles["Heading 1"].font.size = Pt(16)
doc.styles["Heading 2"].font.size = Pt(13)
doc.styles["Heading 3"].font.size = Pt(11.5)


def set_cell_shading(cell, color_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color_hex)
    tcPr.append(shd)


def add_table(headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = ""
        p = hdr_cells[i].paragraphs[0]
        run = p.add_run(h)
        run.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        run.font.size = Pt(10)
        set_cell_shading(hdr_cells[i], "1B2A4A")
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = str(val)
            for p in cells[i].paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9.5)
    if widths:
        for i, w in enumerate(widths):
            for row in table.rows:
                row.cells[i].width = Inches(w)
    doc.add_paragraph()
    return table


def add_bullets(items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def add_numbered(items):
    for item in items:
        doc.add_paragraph(item, style="List Number")


def page_break():
    doc.add_page_break()


# =========================================================
# COVER PAGE
# =========================================================
doc.add_paragraph()
doc.add_paragraph()
doc.add_paragraph()
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run("SPECTRACE AI")
run.font.size = Pt(34)
run.font.bold = True
run.font.color.rgb = NAVY

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = sub.add_run("Business Requirements Document")
run.font.size = Pt(18)
run.font.color.rgb = GREY

sub2 = doc.add_paragraph()
sub2.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = sub2.add_run("Agentic SDLC and Codebase Intelligence Platform")
run.font.size = Pt(13)
run.italic = True
run.font.color.rgb = GREY

doc.add_paragraph()
doc.add_paragraph()
doc.add_paragraph()

meta = doc.add_table(rows=5, cols=2)
meta.alignment = WD_TABLE_ALIGNMENT.CENTER
meta_data = [
    ("Document Type", "Business Requirements Document (BRD)"),
    ("Project Name", "Spectrace AI"),
    ("Version", "1.0"),
    ("Date", datetime.date.today().strftime("%d %B %Y")),
    ("Status", "Draft for Review"),
]
for i, (k, v) in enumerate(meta_data):
    meta.rows[i].cells[0].text = k
    meta.rows[i].cells[0].paragraphs[0].runs[0].font.bold = True
    meta.rows[i].cells[1].text = v
meta.style = "Table Grid"

page_break()

# =========================================================
# DOCUMENT CONTROL
# =========================================================
doc.add_heading("Document Control", level=1)
add_table(
    ["Version", "Date", "Description", "Author"],
    [["1.0", datetime.date.today().strftime("%d %b %Y"), "Initial draft derived from project requirements discussion", "Kartik Singh"]],
    widths=[0.8, 1.0, 3.2, 1.5],
)

doc.add_heading("Table of Contents (Sections)", level=2)
toc_items = [
    "1. Executive Summary",
    "2. Business Objectives",
    "3. Project Scope",
    "4. Stakeholders and User Roles",
    "5. Business Use Case",
    "6. Functional Requirements",
    "7. Agentic AI Workflow Requirements",
    "8. Non-Functional / System Quality Requirements",
    "9. Technical Architecture and Technology Stack",
    "10. Data Model Overview",
    "11. API Surface Overview",
    "12. Application Pages",
    "13. Delivery Plan and Milestones",
    "14. Testing and Evaluation Strategy",
    "15. Acceptance Criteria",
    "16. Assumptions, Constraints and Dependencies",
    "17. Risks",
    "18. Final Deliverables",
]
add_bullets(toc_items)

page_break()

# =========================================================
# 1. EXECUTIVE SUMMARY
# =========================================================
doc.add_heading("1. Executive Summary", level=1)
doc.add_paragraph(
    "Spectrace AI is a full-stack, agentic software-delivery-lifecycle (SDLC) assistant that helps a "
    "software team understand a new requirement before development begins. A user uploads requirement "
    "documents and a source-code project, then submits a change request in natural language. The system "
    "retrieves relevant evidence from both the documents and the codebase using semantic search, proposes "
    "potentially affected files, generates a user story with acceptance criteria and test cases, and "
    "submits the resulting plan for human reviewer approval before it is considered final."
)
doc.add_paragraph(
    "The platform is built around a controlled, stateful AI agent rather than a single unrestrained "
    "language-model call. The agent selects from a fixed set of approved tools, retrieves and grounds its "
    "output in real evidence, asks the user a clarification question when information is missing, and "
    "always pauses for explicit human approval before a plan is finalized. This document defines the "
    "business and functional requirements for the first (mandatory) release of Spectrace AI."
)

doc.add_heading("1.1 Project Snapshot", level=2)
add_table(
    ["Attribute", "Value"],
    [
        ["Project Name", "Spectrace AI"],
        ["Project Type", "Full-stack application with Agentic AI and Retrieval-Augmented Generation (RAG)"],
        ["Frontend", "Next.js with TypeScript"],
        ["Backend", "Python with FastAPI"],
        ["Database", "PostgreSQL (self-managed) with pgvector extension for embeddings"],
        ["Agent Orchestration", "LangGraph"],
        ["Primary Outcome", "A deployed application, an evaluated AI workflow, documented source code, and a demonstration video"],
    ],
    widths=[1.8, 4.5],
)

# =========================================================
# 2. BUSINESS OBJECTIVES
# =========================================================
doc.add_heading("2. Business Objectives", level=1)
add_bullets([
    "Reduce the time a software team spends manually cross-referencing requirement documents against an existing codebase before development starts.",
    "Improve the quality and consistency of user stories, acceptance criteria, and test cases produced at the start of a change request.",
    "Increase confidence in AI-assisted analysis by grounding every generated claim in retrievable, citable evidence rather than unverified model output.",
    "Provide a controlled, auditable human-approval checkpoint so no AI-generated plan is treated as final without reviewer sign-off.",
    "Demonstrate a practical, production-style agentic AI workflow (retrieval, tool selection, state, human-in-the-loop) as a portfolio-grade full-stack build.",
])

# =========================================================
# 3. PROJECT SCOPE
# =========================================================
doc.add_heading("3. Project Scope", level=1)

doc.add_heading("3.1 In Scope (Mandatory, First Release)", level=2)
add_bullets([
    "User registration, authentication (sign in / sign out), and role-based access (Contributor, Reviewer, Administrator).",
    "Project creation, management, and strict data isolation between projects.",
    "Ingestion of requirement documents (PDF, TXT, Markdown) with text extraction, chunking, and embedding generation.",
    "Ingestion of a source-code ZIP archive (TypeScript, JavaScript, Python) with safe-path validation and function/class-level chunking.",
    "Semantic search across requirement and code content, filtered by project and content type.",
    "A stateful agent workflow that classifies the request, selects tools, retrieves evidence, asks clarifying questions when needed, and generates a structured plan.",
    "Human review workflow: approve, edit-and-approve, reject with feedback, or request regeneration.",
    "A visible agent run history / execution timeline.",
    "A repeatable evaluation dataset and measured retrieval/generation quality metrics.",
    "Deployment of the frontend, API, and database to a live, reviewable environment.",
])

doc.add_heading("3.2 Out of Scope (First Release)", level=2)
add_bullets([
    "Automatically modifying source code.",
    "Creating or merging pull requests.",
    "Presenting suggested affected files as guaranteed facts — all suggestions must be labelled as such with a confidence indicator.",
    "Indexing secrets, environment files, build output, dependency folders, or binary files.",
    "Using confidential company code or documents without written authorisation.",
])

doc.add_heading("3.3 Optional Features (Post-Mandatory-Acceptance Only)", level=2)
add_bullets([
    "Hybrid search combining vector similarity with keyword search.",
    "A reranker to improve the ordering of retrieved evidence.",
    "Direct read-only GitHub repository import.",
    "Server-sent events or WebSocket updates for live agent progress.",
    "Export of an approved plan to Markdown or PDF.",
    "Comparison of two versions of a requirement document.",
])

page_break()

# =========================================================
# 4. STAKEHOLDERS
# =========================================================
doc.add_heading("4. Stakeholders and User Roles", level=1)
add_table(
    ["Role", "Allowed Actions"],
    [
        ["Contributor", "Create projects, upload permitted files, submit change requests, and view runs for their own projects."],
        ["Reviewer", "All Contributor actions, plus approve, edit, or reject generated plans."],
        ["Administrator", "Manage users, supported file settings, and project-level access."],
    ],
    widths=[1.5, 4.8],
)

# =========================================================
# 5. BUSINESS USE CASE
# =========================================================
doc.add_heading("5. Business Use Case", level=1)
doc.add_paragraph("Representative scenario submitted by a reviewer:", style="Intense Quote")
q = doc.add_paragraph()
run = q.add_run('"Make mobile OTP verification mandatory before a user can proceed with a loan application."')
run.italic = True

doc.add_paragraph("Expected system behaviour:")
add_numbered([
    "Search uploaded requirement documents and identify the relevant business rule.",
    "Search the indexed codebase for OTP, mobile-verification, and navigation logic.",
    "List potentially affected files with the supporting evidence for each suggestion.",
    "Generate a user story, acceptance criteria, development tasks, and test cases.",
    "Ask a clarification question if the requirement does not state when OTP should be triggered.",
    "Pause and wait for a reviewer to approve, edit, or reject the generated plan.",
])

# =========================================================
# 6. FUNCTIONAL REQUIREMENTS
# =========================================================
doc.add_heading("6. Functional Requirements", level=1)

doc.add_heading("6.1 Authentication and Projects", level=2)
add_bullets([
    "Users can register, sign in, and sign out.",
    "Protected endpoints reject unauthenticated requests.",
    "Users can create, view, rename, and delete projects they are authorised to access.",
    "Project data is isolated so one project cannot retrieve another project's content.",
])

doc.add_heading("6.2 Requirement Document Ingestion", level=2)
add_bullets([
    "Accept PDF, TXT, and Markdown files within a documented size limit.",
    "Extract text and preserve filename, page number, or section information.",
    "Split text into configurable chunks with controlled overlap.",
    "Generate embeddings and store each chunk with project and document metadata.",
    "Display uploaded, processing, ready, or failed status, with an error message on failure.",
    "Allow an authorised user to delete a document and all associated chunks.",
])

doc.add_heading("6.3 Source-Code Ingestion", level=2)
add_bullets([
    "Accept a ZIP archive containing TypeScript, JavaScript, or Python files.",
    "Reject archives with unsafe paths (zip-slip) and enforce file-count and uncompressed-size limits.",
    "Ignore node_modules, virtual environments, build output, generated files, binary files, .env files, and known secret files.",
    "Store repository-relative path, language, symbol name, and available line information.",
    "Chunk code by function or class where practical, with a documented fallback for unsupported syntax.",
])

doc.add_heading("6.4 Retrieval", level=2)
add_bullets([
    "Provide separate requirement and code semantic-search endpoints.",
    "Filter every search by project ID and content type.",
    "Return top results with similarity score and source metadata.",
    "Allow the user to open the evidence used in a generated output.",
    "Use a configurable retrieval limit and minimum relevance threshold.",
])

doc.add_heading("6.5 Analysis and Generated Output", level=2)
doc.add_paragraph("Every completed analysis must include:")
add_bullets([
    "Requirement summary and detected request type.",
    "Clarification questions, when needed.",
    "Relevant requirement evidence with citations.",
    "Potentially affected files with a reason and confidence label.",
    "User story and acceptance criteria.",
    "Frontend, backend, database, and testing tasks as applicable.",
    "Positive, negative, boundary, permission, and regression test cases.",
    "Dependencies, assumptions, and technical risks.",
])

doc.add_heading("6.6 Human Review", level=2)
add_bullets([
    "The generated plan enters a waiting-for-approval state.",
    "A reviewer can approve, edit-and-approve, reject with feedback, or request regeneration.",
    "The system records the reviewer, decision, date, and final content.",
    "Rejected output is never marked approved.",
])

doc.add_heading("6.7 Agent Run History", level=2)
doc.add_paragraph("The interface must display a readable execution timeline, for example:")
timeline = [
    "Analysed change request",
    "Searched 24 requirement chunks",
    "Searched 81 code chunks",
    "Selected 5 evidence items",
    "Generated 12 acceptance and test conditions",
    "Waiting for reviewer approval",
]
for t in timeline:
    p = doc.add_paragraph(t)
    p.paragraph_format.left_indent = Inches(0.3)

page_break()

# =========================================================
# 7. AGENTIC AI WORKFLOW REQUIREMENTS
# =========================================================
doc.add_heading("7. Agentic AI Workflow Requirements", level=1)
doc.add_paragraph(
    "A fixed workflow always runs the same overall steps; within that workflow, the agent chooses its "
    "next action based on the current request and the information already found. The agent uses controlled "
    "tools and clear stopping conditions rather than unrestricted access to the application or operating system."
)

doc.add_heading("7.1 Core Concepts", level=2)
add_table(
    ["Concept", "Use in Spectrace AI"],
    [
        ["Large Language Model", "Understands the request and generates structured outputs."],
        ["Embedding", "Converts requirement text or code into numeric vectors that preserve meaning."],
        ["Semantic Search", "Finds relevant content even when the query and source use different wording."],
        ["RAG", "Retrieves relevant evidence and supplies it to the model before generation."],
        ["Agent", "Chooses which approved tool to call and when to stop or ask a question."],
        ["State", "Stores the current request, retrieved evidence, generated plan, and review status."],
        ["Human Approval", "Prevents the system from finalising a plan without reviewer confirmation."],
    ],
    widths=[1.6, 4.7],
)

doc.add_heading("7.2 Expected Agent Flow", level=2)
add_numbered([
    "Receive and validate the change request.",
    "Classify the request as feature, bug, refactor, security, performance, or documentation.",
    "Decide whether to search requirements, code, existing stories, or more than one source.",
    "Retrieve evidence and assess whether sufficient context exists.",
    "Ask the user a specific clarification question when required information is missing.",
    "Generate the impact analysis, story, acceptance criteria, tasks, and test cases.",
    "Run a review step that checks grounding, completeness, and unsupported claims.",
    "Pause for human approval and resume after approve, edit, or reject feedback.",
])

doc.add_heading("7.3 Agent Tools", level=2)
add_table(
    ["Tool", "Responsibility"],
    [
        ["search_requirements", "Return relevant requirement chunks with document and page metadata."],
        ["search_codebase", "Return relevant code chunks with file, symbol, and line metadata."],
        ["get_file_context", "Retrieve surrounding code for a selected result."],
        ["find_similar_stories", "Find related approved stories or earlier sample requirements."],
        ["request_clarification", "Pause the run and present a focused question to the user."],
        ["generate_plan", "Produce the proposed story, criteria, tasks, risks, and test cases."],
        ["submit_for_approval", "Store the proposal and change the run to waiting-for-review."],
    ],
    widths=[1.8, 4.5],
)

doc.add_heading("7.4 Structured Agent Output", level=2)
doc.add_paragraph(
    "The agent must return validated JSON that the interface can render directly, using a schema validator "
    "(e.g., Pydantic) rather than parsing free-form text. At minimum, the schema must include: summary, "
    "request type, questions, evidence, affected files, user story, acceptance criteria, tasks, test cases, "
    "assumptions, risks, and confidence."
)

page_break()

# =========================================================
# 8. NON-FUNCTIONAL REQUIREMENTS
# =========================================================
doc.add_heading("8. Non-Functional / System Quality Requirements", level=1)
add_table(
    ["Area", "Requirement"],
    [
        ["Security", "Validate file types and archives, isolate projects, protect endpoints, and keep API keys on the server."],
        ["Privacy", "Use sample or authorised data only, and document which external model receives content."],
        ["Reliability", "Handle failed uploads, model errors, and empty retrieval without losing run state."],
        ["Performance", "Record ingestion and analysis duration; show progress for long-running operations."],
        ["Grounding", "Show source evidence and return an insufficient-evidence response instead of inventing facts."],
        ["Observability", "Log run ID, step name, duration, status, and error, without storing secrets."],
        ["Accessibility", "Use labelled controls, keyboard access, visible focus states, and readable contrast."],
    ],
    widths=[1.5, 4.8],
)

# =========================================================
# 9. TECHNICAL ARCHITECTURE
# =========================================================
doc.add_heading("9. Technical Architecture and Technology Stack", level=1)
doc.add_paragraph(
    "The Next.js application presents the user interface and calls the FastAPI service. FastAPI handles "
    "authentication, ingestion, retrieval, and agent execution. A self-managed PostgreSQL instance stores "
    "all application data, with the pgvector extension enabled on the same database to store embeddings "
    "alongside chunk metadata. A managed backend-as-a-service (e.g., Supabase) is intentionally not used "
    "for this project — authentication, migrations, and data access are implemented directly against "
    "PostgreSQL to build backend engineering depth. File bytes are stored outside the database (local disk "
    "in development, object storage in production)."
)
add_table(
    ["Layer", "Technology", "Responsibility"],
    [
        ["Frontend", "Next.js + TypeScript", "Pages, forms, project workspace, run timeline, and approval interface."],
        ["API", "Python + FastAPI", "Validation, authentication, ingestion, retrieval, and application services."],
        ["Agent", "LangGraph", "State, conditional routing, tool calls, interrupts, and resumption."],
        ["Database", "PostgreSQL (self-managed)", "Users, projects, files, runs, steps, plans, and approvals."],
        ["Vector Search", "pgvector (Postgres extension)", "Requirement and code embeddings with metadata filters."],
        ["Model", "Approved LLM provider", "Classification, tool selection, and structured generation."],
        ["Deployment", "Vercel (frontend) + a Python-capable host (API) + managed/self-hosted Postgres", "Deployed frontend, API, and database."],
    ],
    widths=[1.2, 2.0, 3.3],
)

# =========================================================
# 10. DATA MODEL OVERVIEW
# =========================================================
doc.add_heading("10. Data Model Overview", level=1)
add_table(
    ["Table", "Key Information"],
    [
        ["users", "Identity, password hash, role, and timestamps."],
        ["projects", "Name, owner, and timestamps."],
        ["project_members", "Project, user, and permission."],
        ["documents", "Project, filename, type, processing status, and error."],
        ["content_chunks", "Project, document, content type, text, vector, and source metadata."],
        ["change_requests", "Project, request text, classification, and status."],
        ["agent_runs", "Request, state, start time, end time, and result status."],
        ["agent_steps", "Run, tool, input summary, output summary, duration, and status."],
        ["generated_plans", "Structured proposal and version."],
        ["approvals", "Plan, reviewer, decision, feedback, and timestamp."],
        ["evaluation_cases", "Question, expected sources, and expected behaviour."],
        ["evaluation_results", "Retrieved sources, scores, decision, and latency."],
    ],
    widths=[1.8, 4.5],
)

page_break()

# =========================================================
# 11. API SURFACE OVERVIEW
# =========================================================
doc.add_heading("11. API Surface Overview", level=1)
add_table(
    ["Method", "Endpoint", "Purpose"],
    [
        ["POST", "/auth/register", "Create a user."],
        ["POST", "/auth/login", "Authenticate and return a secure session or token."],
        ["GET / POST", "/projects", "List or create projects."],
        ["GET / PATCH / DELETE", "/projects/{id}", "Read, rename, or delete a project."],
        ["POST", "/projects/{id}/documents", "Upload a requirement document."],
        ["POST", "/projects/{id}/codebases", "Upload and validate a source-code archive."],
        ["GET", "/projects/{id}/documents", "List files and processing status."],
        ["POST", "/projects/{id}/search", "Run filtered semantic search."],
        ["POST", "/projects/{id}/requests", "Create a change request."],
        ["POST", "/requests/{id}/analyse", "Start or resume an agent run."],
        ["GET", "/runs/{id}", "Return state, steps, evidence, and generated plan."],
        ["POST", "/runs/{id}/clarification", "Provide requested information and resume."],
        ["POST", "/plans/{id}/decision", "Approve, edit, reject, or request regeneration."],
    ],
    widths=[1.3, 2.5, 2.7],
)

# =========================================================
# 12. APPLICATION PAGES
# =========================================================
doc.add_heading("12. Application Pages", level=1)
add_table(
    ["Page", "Purpose"],
    [
        ["Sign in", "Authenticate the user."],
        ["Dashboard", "List projects and recent analyses."],
        ["Project workspace", "Manage documents, codebase, and change requests."],
        ["Document centre", "Upload files and review processing status."],
        ["Semantic search", "Test requirement and code retrieval independently."],
        ["New analysis", "Enter a change request and start the agent."],
        ["Run details", "Show steps, retrieved evidence, and pending questions."],
        ["Plan review", "Edit, approve, reject, or regenerate the proposal."],
        ["Evaluation", "Run test cases and display measured results."],
    ],
    widths=[1.8, 4.5],
)

page_break()

# =========================================================
# 13. DELIVERY PLAN
# =========================================================
doc.add_heading("13. Delivery Plan and Milestones", level=1)
doc.add_paragraph(
    "Working software should be demonstrable at the end of each phase below. If a phase runs late, "
    "optional features must be removed before mandatory quality is reduced. Dates should be re-baselined "
    "against the project's actual kickoff date if it differs from the reference schedule."
)
add_table(
    ["Milestone", "Evidence of Completion"],
    [
        ["Design and setup", "Wireframes, architecture, schema, repositories, and running frontend and API."],
        ["Application foundation", "Authentication, roles, project CRUD, and protected endpoints."],
        ["Ingestion", "Documents and safe code ZIPs processed with visible status."],
        ["Retrieval", "Embeddings stored and semantic-search results shown with metadata."],
        ["Agent workflow", "Tool selection, persistent state, structured plan, and error handling."],
        ["Human interaction", "Clarification pause, approval interrupt, and successful resumption."],
        ["Quality and evaluation", "Test dataset, automated checks, and recorded metrics."],
        ["Product finish", "Responsive UI, accessibility checks, logs, and failure states."],
        ["Deployment and documentation", "Live URLs, README, API notes, screenshots, and demo video."],
        ["Final demonstration", "End-to-end presentation and explanation of decisions and limitations."],
    ],
    widths=[2.0, 4.3],
)

doc.add_heading("13.1 Daily Working Expectations", level=2)
add_bullets([
    "Push small, understandable commits with meaningful messages.",
    "Maintain a short task board with planned, in-progress, blocked, and completed items.",
    "Record blockers early rather than hiding incomplete work until the final day.",
    "Demonstrate the current build during reviews instead of presenting screenshots alone.",
    "Write tests alongside important backend and agent behaviour.",
])

# =========================================================
# 14. TESTING AND EVALUATION
# =========================================================
doc.add_heading("14. Testing and Evaluation Strategy", level=1)

doc.add_heading("14.1 Required Test Dataset (Minimum 15 Change Requests)", level=2)
add_bullets([
    "Six requests with clear answers and known related files.",
    "Four requests that require evidence from both requirements and code.",
    "Three intentionally ambiguous requests that should trigger clarification.",
    "Two unsupported requests that should produce an insufficient-evidence or escalation response.",
])

doc.add_heading("14.2 Metrics", level=2)
add_table(
    ["Metric", "How to Measure"],
    [
        ["Retrieval hit rate", "Percentage of test cases where an expected source appears in the top results."],
        ["Affected file precision", "Relevant suggested files divided by all suggested files for labelled cases."],
        ["Citation correctness", "Percentage of citations that support the associated statement."],
        ["Clarification accuracy", "Percentage of ambiguous cases where the agent asks an appropriate question."],
        ["Unsupported claim rate", "Percentage of material claims lacking supporting evidence; lower is better."],
        ["Reviewer acceptance", "Percentage of plans accepted without major editing."],
        ["Latency", "Median and slowest duration for retrieval and complete analysis."],
    ],
    widths=[1.8, 4.5],
)

doc.add_heading("14.3 Minimum Test Coverage", level=2)
add_bullets([
    "Authentication and project access tests.",
    "File-validation and unsafe-archive tests.",
    "Document and code chunking tests.",
    "Project-filtered retrieval tests.",
    "Agent routing and invalid tool-input tests.",
    "Clarification pause and resume tests.",
    "Approve, edit, and reject workflow tests.",
    "One complete end-to-end happy path and one failure path.",
])

page_break()

# =========================================================
# 15. ACCEPTANCE CRITERIA
# =========================================================
doc.add_heading("15. Acceptance Criteria", level=1)
add_bullets([
    "A user can sign in, create a project, and upload supported files.",
    "The system indexes requirement text and supported code without exposing ignored secret files.",
    "Semantic search returns project-scoped results with useful source metadata.",
    "The agent uses at least three distinct tools during demonstrated scenarios.",
    "An ambiguous request pauses for clarification and resumes with preserved state.",
    "A generated plan contains cited evidence and affected-file explanations.",
    "The plan cannot become approved without a reviewer decision.",
    "The application handles insufficient evidence without fabricating an answer.",
    "The evaluation report contains repeatable test cases and measured results.",
    "The deployed application, source repository, and documentation are accessible for final review.",
])

# =========================================================
# 16. ASSUMPTIONS / CONSTRAINTS
# =========================================================
doc.add_heading("16. Assumptions, Constraints and Dependencies", level=1)
add_bullets([
    "Only sample, public, or explicitly authorised documents and code will be used — no confidential company data without written authorisation.",
    "A single external LLM provider and embedding provider will be selected and documented in the README, including exactly what data is sent to it.",
    "The project uses a self-managed PostgreSQL database rather than a backend-as-a-service platform, as a deliberate learning constraint.",
    "No model keys, passwords, tokens, or environment files will be committed to source control.",
    "The system is a single-agent, controlled-tool design; multi-agent orchestration is explicitly out of scope for the first release.",
])

# =========================================================
# 17. RISKS
# =========================================================
doc.add_heading("17. Risks", level=1)
add_table(
    ["Risk", "Mitigation"],
    [
        ["Agent enters an unproductive loop between tool calls", "Enforce a maximum step count and explicit stopping conditions in the LangGraph state machine."],
        ["Generated plan includes unsupported claims", "Mandatory grounding/review step before a plan reaches submit_for_approval."],
        ["Code ZIP contains unsafe paths or oversized content", "Archive validation: path traversal checks, file-count limits, and uncompressed-size limits before extraction."],
        ["Retrieval returns irrelevant results for ambiguous queries", "Minimum relevance threshold and insufficient-evidence fallback response."],
        ["Solo delivery timeline slips against the mandatory schedule", "Optional features are deprioritised first; mandatory acceptance criteria are protected."],
    ],
    widths=[2.8, 3.5],
)

# =========================================================
# 18. FINAL DELIVERABLES
# =========================================================
doc.add_heading("18. Final Deliverables", level=1)
add_table(
    ["Deliverable", "Expected Content"],
    [
        ["Source repository", "Frontend, backend, tests, sample data, and useful commit history."],
        ["Deployed application", "Working frontend, API, and database using non-confidential sample data."],
        ["README", "Problem, architecture, setup, environment variables, use cases, and limitations."],
        ["Architecture diagram", "Frontend, API, database, vector search, model, and agent flow."],
        ["API documentation", "Endpoints, request/response examples, and authentication method."],
        ["Database schema", "Tables, relationships, indexes, and vector-field details."],
        ["Evaluation report", "Dataset, metrics, results, observed failures, and improvements."],
        ["Demonstration video", "Three to five minutes covering setup, agent flow, approval, and evaluation."],
    ],
    widths=[1.8, 4.5],
)

doc.add_paragraph()
footer = doc.add_paragraph()
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = footer.add_run("— End of Document —")
run.italic = True
run.font.color.rgb = GREY

doc.save(r"C:\Projects 2026\Spectrace AI\Spectrace_AI_BRD.docx")
print("BRD generated successfully.")
